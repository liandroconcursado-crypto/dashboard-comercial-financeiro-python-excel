"""Criação e validação da planilha Excel profissional."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties

from .config import COLORS, EXCEL_OUTPUT_PATH, FONT_FAMILY, ensure_directories


SHEET_NAMES = [
    "01_Dashboard",
    "02_Indicadores",
    "03_Vendas",
    "04_Despesas",
    "05_DRE_Gerencial",
    "06_Fluxo_Caixa",
    "07_Resumo_Mensal",
]

MONEY_FORMAT = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
MONEY_COMPACT_FORMAT = '"R$" #,##0;[Red]-"R$" #,##0'
PERCENT_FORMAT = "0.0%"
DATE_FORMAT = "dd/mm/yyyy"
INTEGER_FORMAT = "#,##0"

THIN_GRAY = Side(style="thin", color=COLORS["grid"])
LIGHT_BORDER = Border(bottom=Side(style="thin", color="CBD5E1"))


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _font(*, size: int = 10, bold: bool = False, color: str = COLORS["dark_text"]) -> Font:
    return Font(name=FONT_FAMILY, size=size, bold=bold, color=color)


def _python_value(value: Any) -> Any:
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, np.generic):
        return value.item()
    return value


def _setup_sheet(ws, *, gridlines: bool = False) -> None:
    ws.sheet_view.showGridLines = gridlines
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_view.zoomScale = 90


def _title(ws, title: str, subtitle: str, end_column: int, *, dark: bool = False) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    ws.cell(1, 1, title)
    ws.cell(2, 1, subtitle)
    if dark:
        ws.cell(1, 1).font = _font(size=18, bold=True, color=COLORS["white"])
        ws.cell(2, 1).font = _font(size=9, color=COLORS["muted"])
    else:
        ws.cell(1, 1).font = _font(size=18, bold=True, color=COLORS["navy"])
        ws.cell(2, 1).font = _font(size=9, color="64748B")
        ws.cell(2, 1).border = LIGHT_BORDER
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.cell(2, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 22


def _style_header(ws, row: int, start_column: int, end_column: int) -> None:
    for cell in ws.iter_cols(min_col=start_column, max_col=end_column, min_row=row, max_row=row):
        header = cell[0]
        header.fill = _fill(COLORS["navy"])
        header.font = _font(bold=True, color=COLORS["white"])
        header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header.border = Border(left=Side(style="thin", color="FFFFFF"), right=Side(style="thin", color="FFFFFF"))
    ws.row_dimensions[row].height = 30


def _write_dataframe(
    ws,
    frame: pd.DataFrame,
    start_row: int,
    start_col: int,
    table_name: str,
    headers: list[str] | None = None,
) -> tuple[int, int]:
    headers = headers or [str(column) for column in frame.columns]
    for offset, header in enumerate(headers):
        ws.cell(start_row, start_col + offset, header)
    for row_offset, row in enumerate(frame.itertuples(index=False, name=None), start=1):
        for col_offset, value in enumerate(row):
            ws.cell(start_row + row_offset, start_col + col_offset, _python_value(value))
    end_row = start_row + len(frame)
    end_col = start_col + len(frame.columns) - 1
    _style_header(ws, start_row, start_col, end_col)
    table = Table(displayName=table_name, ref=f"{ws.cell(start_row, start_col).coordinate}:{ws.cell(end_row, end_col).coordinate}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    for row in ws.iter_rows(min_row=start_row + 1, max_row=end_row, min_col=start_col, max_col=end_col):
        for cell in row:
            cell.font = _font()
            cell.alignment = Alignment(vertical="center")
    return end_row, end_col


def _build_indicators(ws, analysis: dict[str, Any]) -> dict[str, tuple[int, int, int]]:
    _setup_sheet(ws)
    _title(ws, "Indicadores gerenciais", "Resumo consolidado de janeiro a dezembro de 2025 | dados fictícios", 19)
    kpis = analysis["kpis"]
    indicators = [
        ("Faturamento bruto", kpis["faturamento_total"], MONEY_FORMAT),
        ("Descontos concedidos", kpis["descontos_totais"], MONEY_FORMAT),
        ("Faturamento líquido", kpis["faturamento_liquido"], MONEY_FORMAT),
        ("Custo das mercadorias", kpis["custo_total"], MONEY_FORMAT),
        ("Lucro bruto", kpis["lucro_bruto"], MONEY_FORMAT),
        ("Margem bruta", kpis["margem_media"], PERCENT_FORMAT),
        ("Despesas operacionais", kpis["despesas_totais"], MONEY_FORMAT),
        ("Resultado operacional", kpis["resultado_operacional"], MONEY_FORMAT),
        ("Pedidos", kpis["numero_pedidos"], INTEGER_FORMAT),
        ("Ticket médio", kpis["ticket_medio"], MONEY_FORMAT),
        ("Unidades vendidas", kpis["unidades_vendidas"], INTEGER_FORMAT),
        ("Crescimento do último mês", kpis["crescimento_mensal"], PERCENT_FORMAT),
        ("Produto mais vendido", kpis["produto_mais_vendido"], "@"),
        ("Categoria líder", kpis["categoria_maior_faturamento"], "@"),
        ("Melhor vendedor", kpis["melhor_vendedor"], "@"),
        ("Estado líder", kpis["estado_maior_faturamento"], "@"),
    ]
    indicator_df = pd.DataFrame([(label, value) for label, value, _ in indicators], columns=["Indicador", "Resultado"])
    end_row, _ = _write_dataframe(ws, indicator_df, 4, 1, "tbIndicadores")
    for row, (_, _, number_format) in enumerate(indicators, start=5):
        ws.cell(row, 2).number_format = number_format
        if row in (7, 9, 12):
            ws.cell(row, 2).font = _font(bold=True, color=COLORS["green"] if row != 12 or kpis["resultado_operacional"] >= 0 else COLORS["red"])
    ws.column_dimensions["A"].width = 29
    ws.column_dimensions["B"].width = 24

    ranking_specs = [
        ("produtos", 5, ["Produto", "Unidades", "Faturamento líquido"], ["produto", "unidades_vendidas", "faturamento_liquido"], "tbRankingProdutos"),
        ("categorias", 9, ["Categoria", "Faturamento líquido", "Lucro bruto"], ["categoria", "faturamento_liquido", "lucro_bruto"], "tbRankingCategorias"),
        ("vendedores", 13, ["Vendedor", "Faturamento líquido", "Pedidos"], ["vendedor", "faturamento_liquido", "pedidos"], "tbRankingVendedores"),
        ("estados", 17, ["Estado", "Faturamento líquido", "Pedidos"], ["estado", "faturamento_liquido", "pedidos"], "tbRankingEstados"),
    ]
    ranges: dict[str, tuple[int, int, int]] = {}
    for key, start_col, headers, columns, table_name in ranking_specs:
        frame = analysis["rankings"][key][columns]
        ranking_end, _ = _write_dataframe(ws, frame, 4, start_col, table_name, headers)
        ranges[key] = (start_col, 5, ranking_end)
        for row in range(5, ranking_end + 1):
            if key == "produtos":
                ws.cell(row, start_col + 1).number_format = INTEGER_FORMAT
                ws.cell(row, start_col + 2).number_format = MONEY_FORMAT
            else:
                ws.cell(row, start_col + 1).number_format = MONEY_FORMAT
                ws.cell(row, start_col + 2).number_format = INTEGER_FORMAT if key in ("vendedores", "estados") else MONEY_FORMAT
        ws.column_dimensions[get_column_letter(start_col)].width = 26 if key != "estados" else 12
        ws.column_dimensions[get_column_letter(start_col + 1)].width = 19
        ws.column_dimensions[get_column_letter(start_col + 2)].width = 19
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:B{end_row}"
    return ranges


def _build_sales(ws, vendas: pd.DataFrame) -> None:
    _setup_sheet(ws, gridlines=True)
    _title(ws, "Base de vendas tratada", "Dados validados, padronizados e enriquecidos pelo pipeline Python", 19)
    columns = [
        "data",
        "numero_pedido",
        "cliente",
        "estado",
        "cidade",
        "vendedor",
        "produto",
        "categoria",
        "quantidade",
        "preco_unitario",
        "desconto",
        "custo_unitario",
        "forma_pagamento",
        "faturamento_bruto",
        "valor_desconto",
        "faturamento_liquido",
        "custo",
        "lucro_bruto",
        "margem_percentual",
    ]
    headers = [
        "Data",
        "Número do pedido",
        "Cliente",
        "Estado",
        "Cidade",
        "Vendedor",
        "Produto",
        "Categoria",
        "Quantidade",
        "Preço unitário",
        "Desconto",
        "Custo unitário",
        "Forma de pagamento",
        "Faturamento bruto",
        "Valor do desconto",
        "Faturamento líquido",
        "Custo",
        "Lucro bruto",
        "Margem percentual",
    ]
    end_row, _ = _write_dataframe(ws, vendas[columns], 4, 1, "tbVendas", headers)
    for row in range(5, end_row + 1):
        ws.cell(row, 1).number_format = DATE_FORMAT
        ws.cell(row, 9).number_format = INTEGER_FORMAT
        for col in (10, 12, 14, 15, 16, 17, 18):
            ws.cell(row, col).number_format = MONEY_FORMAT
        for col in (11, 19):
            ws.cell(row, col).number_format = PERCENT_FORMAT
    ws.conditional_formatting.add(
        f"R5:R{end_row}",
        ColorScaleRule(start_type="min", start_color="FCA5A5", mid_type="percentile", mid_value=50, mid_color="FEF3C7", end_type="max", end_color="A7F3D0"),
    )
    ws.conditional_formatting.add(
        f"S5:S{end_row}",
        ColorScaleRule(start_type="min", start_color="FCA5A5", mid_type="percentile", mid_value=50, mid_color="FEF3C7", end_type="max", end_color="A7F3D0"),
    )
    widths = [12, 20, 25, 10, 20, 20, 27, 18, 12, 16, 13, 16, 22, 19, 19, 19, 16, 17, 17]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "C5"


def _build_expenses(ws, despesas: pd.DataFrame) -> None:
    _setup_sheet(ws, gridlines=True)
    _title(ws, "Base de despesas tratada", "Despesas operacionais padronizadas para consolidação gerencial", 5)
    headers = ["Data", "Categoria", "Descrição", "Tipo", "Valor"]
    end_row, _ = _write_dataframe(ws, despesas[["data", "categoria", "descricao", "tipo", "valor"]], 4, 1, "tbDespesas", headers)
    for row in range(5, end_row + 1):
        ws.cell(row, 1).number_format = DATE_FORMAT
        ws.cell(row, 5).number_format = MONEY_FORMAT
    for letter, width in {"A": 13, "B": 28, "C": 34, "D": 14, "E": 18}.items():
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A5"


def _build_dre(ws, dre: pd.DataFrame) -> None:
    _setup_sheet(ws)
    _title(ws, "DRE gerencial simplificada", "Visão demonstrativa, sem finalidade fiscal ou contábil oficial", 14)
    ws.cell(4, 1, "Linha gerencial")
    for index, month in enumerate(dre["mes_label"], start=2):
        ws.cell(4, index, month)
    ws.cell(4, 14, "Total")
    _style_header(ws, 4, 1, 14)
    rows = [
        ("Receita Bruta", "receita_bruta", 1),
        ("(-) Descontos", "descontos", -1),
        ("= Receita Líquida", "receita_liquida", 1),
        ("(-) Custo das Mercadorias", "custo_mercadorias", -1),
        ("= Lucro Bruto", "lucro_bruto", 1),
        ("(-) Despesas Operacionais", "despesas_operacionais", -1),
        ("= Resultado Operacional", "resultado_operacional", 1),
    ]
    for row_index, (label, column, sign) in enumerate(rows, start=5):
        label_cell = ws.cell(row_index, 1, label)
        if label.startswith("="):
            label_cell.data_type = "s"
        label_cell.font = _font(bold=label.startswith("="))
        for col_index, value in enumerate(dre[column], start=2):
            ws.cell(row_index, col_index, float(value) * sign)
            ws.cell(row_index, col_index).number_format = MONEY_FORMAT
        ws.cell(row_index, 14, f"=SUM(B{row_index}:M{row_index})")
        ws.cell(row_index, 14).number_format = MONEY_FORMAT
        if label.startswith("="):
            for cell in ws[row_index][0:14]:
                cell.fill = _fill("E2E8F0")
                cell.font = _font(bold=True, color=COLORS["navy"])
        if label == "= Resultado Operacional":
            for cell in ws[row_index][0:14]:
                cell.fill = _fill("D1FAE5")
                cell.font = _font(bold=True, color="065F46")
    ws.conditional_formatting.add("B11:N11", CellIsRule(operator="lessThan", formula=["0"], fill=_fill("FEE2E2"), font=Font(color="B91C1C", bold=True)))
    ws.conditional_formatting.add("B11:N11", CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=_fill("D1FAE5"), font=Font(color="065F46", bold=True)))
    ws.merge_cells("A14:N14")
    ws["A14"] = "Esta DRE é gerencial, simplificada e exclusivamente demonstrativa. Não substitui demonstrações contábeis ou fiscais oficiais."
    ws["A14"].font = _font(size=9, color="64748B")
    ws["A14"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[14].height = 34
    ws.column_dimensions["A"].width = 32
    for column in range(2, 14):
        ws.column_dimensions[get_column_letter(column)].width = 16
    ws.column_dimensions["N"].width = 19
    ws.freeze_panes = "B5"


def _build_cash_flow(ws, cash: pd.DataFrame) -> None:
    _setup_sheet(ws)
    _title(ws, "Fluxo de caixa mensal simplificado", "Entradas menos saídas de CMV e despesas operacionais", 19)
    headers = ["Mês", "Entradas", "Saídas — CMV", "Saídas — despesas", "Saídas totais", "Saldo mensal", "Saldo acumulado"]
    ws.append([])
    frame = cash[["mes_label", "entradas", "saidas_cmv", "saidas_despesas", "saidas_totais", "saldo_mensal", "saldo_acumulado"]].copy()
    end_row, _ = _write_dataframe(ws, frame, 4, 1, "tbFluxoCaixa", headers)
    for row in range(5, end_row + 1):
        for col in range(2, 8):
            ws.cell(row, col).number_format = MONEY_FORMAT
        if row == 5:
            ws.cell(row, 7, "=F5")
        else:
            ws.cell(row, 7, f"=G{row - 1}+F{row}")
    for column, width in {"A": 13, "B": 18, "C": 18, "D": 20, "E": 18, "F": 18, "G": 20}.items():
        ws.column_dimensions[column].width = width
    ws.conditional_formatting.add(f"F5:G{end_row}", ColorScaleRule(start_type="min", start_color="FCA5A5", mid_type="num", mid_value=0, mid_color="FEF3C7", end_type="max", end_color="A7F3D0"))
    ws.freeze_panes = "A5"

    chart = LineChart()
    chart.title = "Entradas, saídas e saldo mensal"
    chart.style = 13
    chart.height = 8.2
    chart.width = 18.5
    categories = Reference(ws, min_col=1, min_row=5, max_row=end_row)
    series_specs = [(2, "Entradas", COLORS["cyan"]), (5, "Saídas totais", COLORS["orange"]), (6, "Saldo mensal", COLORS["green"])]
    for column, _, _ in series_specs:
        chart.add_data(Reference(ws, min_col=column, min_row=4, max_row=end_row), titles_from_data=True)
    chart.set_categories(categories)
    chart.y_axis.title = "R$"
    chart.x_axis.title = "Mês"
    chart.legend.position = "b"
    for series, (_, title, color) in zip(chart.series, series_specs, strict=True):
        series.tx = SeriesLabel(v=title)
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 22000
    ws.add_chart(chart, "I4")


def _build_monthly(ws, monthly: pd.DataFrame) -> None:
    _setup_sheet(ws, gridlines=True)
    _title(ws, "Resumo mensal", "Consolidação temporal de vendas, custos, despesas e resultado", 14)
    columns = [
        "mes_data",
        "mes_label",
        "faturamento_bruto",
        "descontos",
        "faturamento_liquido",
        "custo",
        "lucro_bruto",
        "despesas",
        "resultado_operacional",
        "margem_percentual",
        "pedidos",
        "ticket_medio",
        "unidades_vendidas",
        "crescimento_mensal",
    ]
    headers = [
        "Data do mês",
        "Mês",
        "Faturamento bruto",
        "Descontos",
        "Faturamento líquido",
        "Custo",
        "Lucro bruto",
        "Despesas",
        "Resultado operacional",
        "Margem bruta",
        "Pedidos",
        "Ticket médio",
        "Unidades vendidas",
        "Crescimento mensal",
    ]
    end_row, _ = _write_dataframe(ws, monthly[columns], 4, 1, "tbResumoMensal", headers)
    for row in range(5, end_row + 1):
        ws.cell(row, 1).number_format = "mmm/yyyy"
        for col in range(3, 10):
            ws.cell(row, col).number_format = MONEY_FORMAT
        ws.cell(row, 10).number_format = PERCENT_FORMAT
        ws.cell(row, 11).number_format = INTEGER_FORMAT
        ws.cell(row, 12).number_format = MONEY_FORMAT
        ws.cell(row, 13).number_format = INTEGER_FORMAT
        ws.cell(row, 14).number_format = PERCENT_FORMAT
    ws.conditional_formatting.add(f"I5:I{end_row}", ColorScaleRule(start_type="min", start_color="FCA5A5", mid_type="num", mid_value=0, mid_color="FEF3C7", end_type="max", end_color="A7F3D0"))
    ws.conditional_formatting.add(f"N5:N{end_row}", ColorScaleRule(start_type="min", start_color="FCA5A5", mid_type="num", mid_value=0, mid_color="FEF3C7", end_type="max", end_color="A7F3D0"))
    widths = [14, 12, 20, 16, 20, 16, 18, 17, 22, 16, 12, 17, 18, 19]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "C5"


def _add_kpi_card(ws, start_col: int, label: str, value: Any, number_format: str, accent: str) -> None:
    end_col = start_col + 2
    for row in range(5, 9):
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row, col)
            cell.fill = _fill(COLORS["navy_light"])
            cell.border = Border(top=THIN_GRAY if row == 5 else None, bottom=THIN_GRAY if row == 8 else None, left=THIN_GRAY if col == start_col else None, right=THIN_GRAY if col == end_col else None)
    ws.merge_cells(start_row=5, start_column=start_col, end_row=5, end_column=end_col)
    ws.merge_cells(start_row=6, start_column=start_col, end_row=8, end_column=end_col)
    label_cell = ws.cell(5, start_col)
    value_cell = ws.cell(6, start_col)
    label_cell.value = label.upper()
    value_cell.value = value
    label_cell.font = _font(size=9, bold=True, color=COLORS["muted"])
    value_cell.font = _font(size=17, bold=True, color=COLORS["white"])
    label_cell.alignment = Alignment(horizontal="left", vertical="center")
    value_cell.alignment = Alignment(horizontal="left", vertical="center")
    value_cell.number_format = number_format
    for col in range(start_col, end_col + 1):
        ws.cell(5, col).border = Border(top=Side(style="medium", color=accent), bottom=THIN_GRAY, left=THIN_GRAY if col == start_col else None, right=THIN_GRAY if col == end_col else None)


def _line_chart(ws_source, min_col: int, min_row: int, max_row: int, title: str, color: str) -> LineChart:
    chart = LineChart()
    chart.title = title
    chart.style = 13
    chart.height = 7.4
    chart.width = 13.2
    chart.add_data(Reference(ws_source, min_col=min_col, min_row=min_row - 1, max_row=max_row), titles_from_data=True)
    chart.set_categories(Reference(ws_source, min_col=2, min_row=min_row, max_row=max_row))
    chart.legend = None
    chart.y_axis.title = "R$"
    chart.x_axis.title = "Mês"
    chart.display_blanks = "gap"
    if chart.series:
        series = chart.series[0]
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 26000
        series.marker.symbol = "circle"
        series.marker.size = 5
    return chart


def _bar_chart(ws_source, label_col: int, value_col: int, min_row: int, max_row: int, title: str, color: str, *, currency: bool = True) -> BarChart:
    chart = BarChart()
    chart.type = "bar"
    chart.style = 13
    chart.title = title
    chart.height = 7.4
    chart.width = 13.2
    chart.add_data(Reference(ws_source, min_col=value_col, min_row=min_row - 1, max_row=max_row), titles_from_data=True)
    chart.set_categories(Reference(ws_source, min_col=label_col, min_row=min_row, max_row=max_row))
    chart.legend = None
    chart.y_axis.title = ""
    chart.x_axis.title = "R$" if currency else "Unidades"
    if chart.series:
        chart.series[0].graphicalProperties.solidFill = color
        chart.series[0].graphicalProperties.line.solidFill = color
    return chart


def _build_dashboard(ws, analysis: dict[str, Any], indicators_ws, monthly_ws, ranges: dict[str, tuple[int, int, int]]) -> None:
    _setup_sheet(ws)
    for row in ws.iter_rows(min_row=1, max_row=58, min_col=1, max_col=19):
        for cell in row:
            cell.fill = _fill(COLORS["navy"])
    _title(ws, "Dashboard Comercial e Financeiro", "Visão executiva | janeiro a dezembro de 2025 | dados fictícios", 19, dark=True)
    ws["A3"] = "Atualização automática a partir das bases tratadas pelo pipeline Python"
    ws["A3"].font = _font(size=9, color=COLORS["muted"])
    kpis = analysis["kpis"]
    cards = [
        (2, "Faturamento líquido", kpis["faturamento_liquido"], MONEY_COMPACT_FORMAT, COLORS["cyan"]),
        (5, "Lucro bruto", kpis["lucro_bruto"], MONEY_COMPACT_FORMAT, COLORS["green"]),
        (8, "Margem bruta", kpis["margem_media"], PERCENT_FORMAT, COLORS["blue"]),
        (11, "Despesas", kpis["despesas_totais"], MONEY_COMPACT_FORMAT, COLORS["orange"]),
        (14, "Pedidos", kpis["numero_pedidos"], INTEGER_FORMAT, COLORS["cyan"]),
        (17, "Ticket médio", kpis["ticket_medio"], MONEY_COMPACT_FORMAT, COLORS["green"]),
    ]
    for start_col, label, value, fmt, accent in cards:
        _add_kpi_card(ws, start_col, label, value, fmt, accent)
    for col in range(1, 20):
        ws.column_dimensions[get_column_letter(col)].width = 10.5
    for row in range(1, 59):
        ws.row_dimensions[row].height = 20

    monthly_end = 4 + len(analysis["mensal"])
    ws.add_chart(_line_chart(monthly_ws, 5, 5, monthly_end, "Faturamento líquido mensal", COLORS["cyan"]), "A10")
    ws.add_chart(_line_chart(monthly_ws, 7, 5, monthly_end, "Lucro bruto mensal", COLORS["green"]), "J10")

    product_col, product_start, product_end = ranges["produtos"]
    category_col, category_start, category_end = ranges["categorias"]
    seller_col, seller_start, seller_end = ranges["vendedores"]
    state_col, state_start, state_end = ranges["estados"]
    ws.add_chart(_bar_chart(indicators_ws, category_col, category_col + 1, category_start, category_end, "Faturamento por categoria", COLORS["blue"]), "A26")
    ws.add_chart(_bar_chart(indicators_ws, seller_col, seller_col + 1, seller_start, seller_end, "Faturamento por vendedor", COLORS["green"]), "J26")
    ws.add_chart(_bar_chart(indicators_ws, product_col, product_col + 1, product_start, min(product_start + 9, product_end), "Top 10 produtos", COLORS["orange"], currency=False), "A42")
    ws.add_chart(_bar_chart(indicators_ws, state_col, state_col + 1, state_start, state_end, "Distribuição regional", COLORS["cyan"]), "J42")
    ws.freeze_panes = "A4"
    ws.sheet_view.zoomScale = 75


def create_excel_report(
    vendas: pd.DataFrame,
    despesas: pd.DataFrame,
    analysis: dict[str, Any],
    output_path: Path = EXCEL_OUTPUT_PATH,
) -> Path:
    """Cria o arquivo Excel final com sete abas na ordem solicitada."""
    ensure_directories()
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    wb.properties.title = "Dashboard Comercial e Financeiro"
    wb.properties.subject = "Automação de Excel com Python — projeto demonstrativo"
    wb.properties.creator = "Projeto demonstrativo de portfólio"
    wb.properties.created = datetime(2026, 9, 5)
    sheets = {name: wb.create_sheet(name) for name in SHEET_NAMES}

    ranges = _build_indicators(sheets["02_Indicadores"], analysis)
    _build_sales(sheets["03_Vendas"], vendas)
    _build_expenses(sheets["04_Despesas"], despesas)
    _build_dre(sheets["05_DRE_Gerencial"], analysis["dre"])
    _build_cash_flow(sheets["06_Fluxo_Caixa"], analysis["fluxo_caixa"])
    _build_monthly(sheets["07_Resumo_Mensal"], analysis["mensal"])
    _build_dashboard(sheets["01_Dashboard"], analysis, sheets["02_Indicadores"], sheets["07_Resumo_Mensal"], ranges)

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = COLORS["cyan"] if ws.title in ("01_Dashboard", "02_Indicadores") else COLORS["navy"]
        ws.sheet_view.selection[0].activeCell = "A1"
        ws.sheet_view.selection[0].sqref = "A1"
    wb.active = 0
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def validate_workbook(path: Path) -> dict[str, Any]:
    """Reabre o XLSX e verifica estrutura, tabelas, gráficos e fórmulas."""
    if not path.exists() or path.stat().st_size < 50_000:
        raise ValueError("O arquivo Excel não existe ou está inesperadamente pequeno.")
    wb = load_workbook(path, data_only=False)
    if wb.sheetnames != SHEET_NAMES:
        raise ValueError(f"Ordem de abas incorreta: {wb.sheetnames}")
    if len(wb["01_Dashboard"]._charts) != 6:
        raise ValueError("O dashboard deve conter exatamente seis gráficos.")
    if len(wb["03_Vendas"].tables) != 1 or len(wb["04_Despesas"].tables) != 1:
        raise ValueError("As bases de vendas e despesas devem estar formatadas como tabelas.")
    if wb["03_Vendas"].max_row < 1505:
        raise ValueError("A aba de vendas contém menos de 1.500 registros tratados.")
    error_tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")
    formula_count = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formula_count += 1
                    if any(token in str(cell.value) for token in error_tokens):
                        raise ValueError(f"Fórmula inválida em {ws.title}!{cell.coordinate}: {cell.value}")
    result = {
        "abas": wb.sheetnames,
        "graficos_dashboard": len(wb["01_Dashboard"]._charts),
        "linhas_vendas": wb["03_Vendas"].max_row - 4,
        "linhas_despesas": wb["04_Despesas"].max_row - 4,
        "formulas": formula_count,
        "tamanho_bytes": path.stat().st_size,
    }
    wb.close()
    return result
